from pptx import Presentation
from pptx.util import Emu
import sys
sys.path.insert(0,'.')
from lib import est_lines
def inch(v): return v/914400
prs=Presentation('KD-PAY-Airwallex-Partnership.pptx')
issues=[]
for i,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        # off-slide check
        if sh.left is None: continue
        r=inch(sh.left+sh.width); b=inch(sh.top+sh.height)
        if r>13.35 or b>7.52 or inch(sh.left)<-0.01 or inch(sh.top)<-0.01:
            issues.append((i,'off-slide',sh.shape_type,inch(sh.left),inch(sh.top),round(r,2),round(b,2)))
        if not sh.has_text_frame: continue
        tf=sh.text_frame
        h=inch(sh.height); w=inch(sh.width)
        total=0
        for p in tf.paragraphs:
            if not p.runs: 
                total+=0.15; continue
            size=p.runs[0].font.size.pt if p.runs[0].font.size else 14
            bold=bool(p.runs[0].font.bold)
            txt="".join(r.text for r in p.runs)
            ls=p.line_spacing if isinstance(p.line_spacing,float) else 1.15
            n=est_lines(txt,w-0.02,size,bold)
            total+=n*size*1.2*ls/72
            if p.space_after: total+=p.space_after.pt/72
        if total>h+0.06:
            issues.append((i,'text taller than box',tf.text[:48].replace('\n',' '),round(total,2),round(h,2)))
print(len(issues),'issue(s)')
for x in issues[:25]: print(' ',x[0],x[1],'|',x[2],x[3:])
