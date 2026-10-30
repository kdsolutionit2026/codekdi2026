"""Small layout toolkit over python-pptx: one visual language for the whole deck."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree

# ---------- palette ----------
INK      = RGBColor(0x0A, 0x1A, 0x2B)
INK_2    = RGBColor(0x12, 0x2C, 0x44)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
TEXT     = RGBColor(0x10, 0x22, 0x36)
MUTED    = RGBColor(0x5B, 0x6E, 0x82)
LINE     = RGBColor(0xDF, 0xE6, 0xEE)
SURFACE  = RGBColor(0xF4, 0xF7, 0xFA)
ACCENT   = RGBColor(0x00, 0xB5, 0x89)
ACCENT_D = RGBColor(0x00, 0x8A, 0x69)
ACCENT_S = RGBColor(0xE7, 0xF7, 0xF2)
BLUE     = RGBColor(0x2F, 0x6F, 0xED)
BLUE_S   = RGBColor(0xE8, 0xF0, 0xFE)
AMBER    = RGBColor(0xE8, 0xA3, 0x1C)
DARK_MUT = RGBColor(0x9F, 0xB3, 0xC6)
DARK_LN  = RGBColor(0x21, 0x3B, 0x55)


# ---------- metrics (conservative estimates; Arial is narrower than DejaVu) ----------
def est_lines(txt, width_in, size_pt, bold=False):
    """Estimate wrapped line count. Explicit \n always starts a new line."""
    if txt is None:
        return 0
    per_char = (0.520 if bold else 0.485) * size_pt / 72.0   # inches per character
    cap = max(6, int(width_in / per_char))
    n = 0
    for seg in str(txt).split("\n"):
        words, cur = seg.split(" "), ""
        lines = 1
        for w in words:
            trial = (cur + " " + w).strip()
            if len(trial) <= cap or not cur:
                cur = trial
            else:
                lines += 1
                cur = w
        n += lines
    return n

def text_h(txt, width_in, size_pt, bold=False, spacing=1.15):
    return est_lines(txt, width_in, size_pt, bold) * size_pt * 1.20 * spacing / 72.0

FONT = "Arial"
W, H = 13.333, 7.5
ML, MR = 0.72, 0.72
CW = W - ML - MR          # content width 11.893

def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    return prs

def slide(prs, bg=WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    f = s.background.fill
    f.solid(); f.fore_color.rgb = bg
    return s

# ---------- text primitives ----------
def _rpr(run, size, color, bold, italic, spc=None):
    f = run.font
    f.name = FONT; f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = color
    if spc is not None:
        run.font._rPr.set("spc", str(int(spc * 100)))

def text(s, x, y, w, h, value, size=14, color=TEXT, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.12, space_after=0,
         caps=False, tracking=None, margins=(0, 0, 0, 0)):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    tf.vertical_anchor = anchor
    lines = value.split("\n") if isinstance(value, str) else list(value)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = ln.upper() if caps else ln
        _rpr(r, size, color, bold, italic, tracking)
    return box

def rich(s, x, y, w, h, spans, size=11, spacing=1.2, align=PP_ALIGN.LEFT):
    """One paragraph of mixed runs: spans = [(text, bold, color), ...]."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    for txt, bold, color in spans:
        r = p.add_run()
        r.text = txt
        _rpr(r, size, color, bold, False)
    return box


def para_list(s, x, y, w, h, items, size=12.5, color=TEXT, bullet_color=ACCENT,
              gap=7, spacing=1.16, bullet="—", indent=0.24):
    """items: list of str, or (bold_lead, rest) tuples."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        p.space_after = Pt(gap)
        pPr = p._pPr if p._pPr is not None else p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(indent))))
        pPr.set("indent", str(-Emu(Inches(indent))))
        bu = etree.SubElement(pPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}buClr")
        srgb = etree.SubElement(bu, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
        srgb.set("val", "%02X%02X%02X" % (bullet_color[0], bullet_color[1], bullet_color[2]))
        bc = etree.SubElement(pPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}buFont")
        bc.set("typeface", FONT)
        bch = etree.SubElement(pPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar")
        bch.set("char", bullet)
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run(); r1.text = lead; _rpr(r1, size, color, True, False)
            if rest:
                r2 = p.add_run(); r2.text = rest; _rpr(r2, size, color, False, False)
        else:
            r = p.add_run(); r.text = item; _rpr(r, size, color, False, False)
    return box

def rect(s, x, y, w, h, fill=None, line=None, radius=0.0, line_w=0.75, shape=None):
    shp = s.shapes.add_shape(shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE),
                             Inches(x), Inches(y), Inches(w), Inches(h))
    if radius and shp.adjustments:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    shp.text_frame.word_wrap = True
    return shp

def hline(s, x, y, w, color=LINE, weight=0.75):
    ln = s.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color; ln.line.width = Pt(weight)
    return ln

def vline(s, x, y, h, color=LINE, weight=0.75):
    ln = s.shapes.add_connector(1, Inches(x), Inches(y), Inches(x), Inches(y + h))
    ln.line.color.rgb = color; ln.line.width = Pt(weight)
    return ln

def arrow(s, x, y, w, color=ACCENT, weight=1.5):
    ln = s.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color; ln.line.width = Pt(weight)
    ln.line._get_or_add_ln().set("cap", "rnd")
    tail = etree.SubElement(ln.line._get_or_add_ln(),
        "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    tail.set("type", "triangle")
    return ln

# ---------- deck furniture ----------
def header(s, kicker, title, dark=False, title_size=30, lead=None, lead_w=None):
    tcol = WHITE if dark else TEXT
    text(s, ML, 0.52, CW, 0.26, kicker, size=9.5, color=ACCENT, bold=True, caps=True, tracking=1.6)
    t_lines = est_lines(title, CW, title_size, bold=True)
    text(s, ML, 0.82, CW, t_lines * title_size * 1.08 / 72 + 0.12, title, size=title_size,
         color=tcol, bold=True, spacing=1.02)
    y = 0.82 + t_lines * title_size * 1.20 * 1.02 / 72 + 0.10
    if lead:
        lw = lead_w or min(CW, 10.2)
        text(s, ML, y, lw, text_h(lead, lw, 13.5) + 0.1, lead, size=13.5,
             color=DARK_MUT if dark else MUTED, spacing=1.25)
        return y + text_h(lead, lw, 13.5) + 0.30
    return y + 0.22


def footer(s, n, dark=False, label="KD PAY LLC  ·  Airwallex partnership proposal  ·  Confidential"):
    col = DARK_MUT if dark else MUTED
    hline(s, ML, 6.92, CW, color=DARK_LN if dark else LINE, weight=0.6)
    text(s, ML, 7.02, 9.0, 0.24, label, size=8, color=col)
    text(s, W - MR - 0.8, 7.02, 0.8, 0.24, f"{n:02d}", size=8.5, color=col, bold=True, align=PP_ALIGN.RIGHT)

def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt

# ---------- components ----------
def stat(s, x, y, w, value, label, sub=None, color=TEXT, vsize=30, dark=False):
    text(s, x, y, w, 0.62, value, size=vsize, bold=True, color=color, spacing=1.0)
    text(s, x, y + 0.60, w, 0.30, label, size=10, color=DARK_MUT if dark else MUTED,
         caps=True, bold=True, tracking=1.1)
    if sub:
        text(s, x, y + 0.86, w, 0.40, sub, size=9.5, color=DARK_MUT if dark else MUTED, spacing=1.15)

def card(s, x, y, w, h, num, title, body, dark=False, accent=ACCENT, body_size=11):
    rect(s, x, y, w, h, fill=INK_2 if dark else SURFACE, line=None, radius=0.06)
    ty = y + 0.22
    if num:
        text(s, x + 0.26, ty, 1.2, 0.22, num, size=9, color=accent, bold=True, tracking=1.4)
        ty += 0.26
    tl = est_lines(title, w - 0.52, 13.5, bold=True)
    text(s, x + 0.26, ty, w - 0.52, tl * 13.5 * 1.18 / 72 + 0.08, title, size=13.5, bold=True,
         color=WHITE if dark else TEXT, spacing=1.05)
    by = ty + tl * 13.5 * 1.20 * 1.05 / 72 + 0.18
    avail = y + h - by - 0.18
    text(s, x + 0.26, by, w - 0.52, avail, body, size=body_size,
         color=DARK_MUT if dark else MUTED, spacing=1.22)
    if text_h(body, w - 0.52, body_size, spacing=1.22) > avail:
        print(f"  !! card overflow: {title[:40]!r} needs {text_h(body, w-0.52, body_size, spacing=1.22):.2f}in, has {avail:.2f}in")


def table(s, x, y, w, cols, rows, dark=False, header_size=9.5, body_size=10.5, row_h=0.52,
          head_h=0.34, zebra=True, col_colors=None, bold_first=True):
    """cols: [(label, width_fraction)], rows: list of tuples of strings."""
    tcol = WHITE if dark else TEXT
    mcol = DARK_MUT if dark else MUTED
    cx = x
    if cols[0][0]:
        for label, frac in cols:
            if label:
                text(s, cx, y, w * frac - 0.12, head_h, label, size=header_size, color=mcol,
                     bold=True, caps=True, tracking=0.9)
            cx += w * frac
        hline(s, x, y + head_h - 0.02, w, color=DARK_LN if dark else LINE, weight=0.9)
    ry = y + head_h + 0.02
    for ri, row in enumerate(rows):
        if zebra and ri % 2 == 1:
            rect(s, x - 0.08, ry - 0.04, w + 0.16, row_h, fill=INK_2 if dark else SURFACE, radius=0.03)
        cx = x
        for ci, (val, (label, frac)) in enumerate(zip(row, cols)):
            colr = tcol
            if col_colors and ci in col_colors:
                colr = col_colors[ci]
            bold = ci == 0 and bold_first
            text(s, cx, ry + 0.12, w * frac - 0.16, row_h - 0.16, val, size=body_size,
                 color=colr, bold=bold, spacing=1.12)
            cx += w * frac
        ry += row_h
    return ry
